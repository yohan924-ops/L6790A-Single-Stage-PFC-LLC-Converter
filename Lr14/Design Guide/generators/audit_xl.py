# -*- coding: utf-8 -*-
"""Structural audit of the ST workbook - find what nobody looked at.

C.12 was a pole computed in D104 and referenced by nothing. C.15 was a selected
value 2.6x above its own limit with no warning. Both are structural: a machine
can find them, a reader will not. So this looks for the SHAPES of those faults
across every sheet rather than re-checking arithmetic:

  1. dead cells     - a formula whose result no cell ever reads
  2. errors         - #REF!, #DIV/0!, #VALUE!, #NAME? sitting in the cache
  3. dangling refs  - a formula that reads a cell holding nothing
  4. limit breaches - a "Selected Value" (column F) outside the limit the
                      calculated column next to it computes
  5. constants      - a literal inside a formula that also exists as a cell,
                      i.e. a value that will not follow when the spec changes
  6. label/formula  - a row named "...max" built with MIN(), and the reverse

Nothing here is a verdict on its own; every hit is a place to look.

    python audit_xl.py            summary + findings
    python audit_xl.py -v         include the low-signal categories
"""
import io
import os
import re
import sys
from collections import defaultdict

import openpyxl


def require_readable(path):
    """Excel takes an exclusive lock while the workbook is open. Say so
    plainly rather than dying inside openpyxl with a bare PermissionError."""
    try:
        open(path, 'rb').close()
    except PermissionError:
        print('워크북이 Excel 에서 열려 있어 읽을 수 없다. 닫고 다시 실행할 것:')
        print('  ' + path)
        raise SystemExit(2)


HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
XL = os.path.join(ROOT, 'Calculation Excel Sheet',
                  'L6790_spreadsheet_r1_0_corrected_rev0_6.xlsx')
if os.environ.get('XL_OVERRIDE'):          # a variant instead of the canonical workbook
    XL = os.environ['XL_OVERRIDE']

# sheets whose job is bulk tabulation, not design - dead cells there are normal
BULK = {'RTC', 'CC', 'OutCap_Sel_DB', 'CHANGELOG_rev0_6', 'BOM&Schematics'}
ERRS = ('#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#NUM!', '#NULL!')

REF = re.compile(r"(?:'([^']+)'|([A-Za-z][A-Za-z0-9_. &]*))?!?\$?([A-Z]{1,2})\$?(\d+)")
# a reference must not sit inside a longer identifier: LOG10 is not cell OG10
CELLREF = re.compile(r"(?<![A-Za-z0-9_.$])(?:('[^']+'|[A-Za-z][A-Za-z0-9_.&]*)!)?"
                     r"(\$?[A-Z]{1,2}\$?\d+)(?![A-Za-z0-9_(])")


def norm(sheet, ref):
    return '%s!%s' % (sheet, ref.replace('$', ''))


def main():
    verbose = '-v' in sys.argv
    require_readable(XL)
    wf = openpyxl.load_workbook(XL)                 # formulas
    wv = openpyxl.load_workbook(XL, data_only=True)  # cached values

    formulas = {}          # norm ref -> formula text
    values = {}            # norm ref -> cached value
    labels = {}            # norm ref -> the label in column C of the same row
    read_by = defaultdict(set)

    for ws in wf.worksheets:
        wsv = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                key = norm(ws.title, c.coordinate)
                values[key] = wsv[c.coordinate].value
                if isinstance(c.value, str) and c.value.startswith('='):
                    formulas[key] = c.value
        for row in ws.iter_rows(min_col=3, max_col=3):
            for c in row:
                if isinstance(c.value, str):
                    for col in 'DEFGHI':
                        labels[norm(ws.title, '%s%d' % (col, c.row))] = c.value

    # who reads whom
    for key, f in formulas.items():
        sheet = key.split('!')[0]
        for m in CELLREF.finditer(f):
            tgt_sheet = (m.group(1) or sheet).strip("'")
            read_by[norm(tgt_sheet, m.group(2))].add(key)

    hits = defaultdict(list)

    # ---------------------------------------------------------- 1. dead cells
    for key, f in formulas.items():
        sheet = key.split('!')[0]
        if sheet in BULK or read_by[key]:
            continue
        col = re.match(r'[^!]+!([A-Z]+)', key).group(1)
        if col not in 'DF':                        # E is the unit column
            continue
        if re.fullmatch(r'="[^"]*"', f) or 'IF(' in f and '""' in f:
            continue                               # decorative / message cells
        if re.fullmatch(r"='[^']+'!\$?[A-Z]+\$?\d+", f) or re.fullmatch(r'=[A-Z]+\d+', f):
            continue                               # a plain echo of another cell
        hits['1 계산만 하고 아무도 안 읽는 셀'].append(
            '%-34s %-26.26s = %s' % (key, labels.get(key, ''), f[:56]))

    # ------------------------------------------------------------- 2. errors
    for key, v in values.items():
        if isinstance(v, str) and v in ERRS:
            hits['2 캐시에 오류값'].append('%-34s %-26.26s %s'
                                       % (key, labels.get(key, ''), v))

    # ------------------------------------------------------- 3. dangling refs
    for key, f in formulas.items():
        if key.split('!')[0] in BULK:
            continue
        sheet = key.split('!')[0]
        for m in CELLREF.finditer(f):
            tgt = norm((m.group(1) or sheet).strip("'"), m.group(2))
            if (tgt not in values and tgt.split('!')[0] not in BULK
                    and tgt.split('!')[0] in {w.title for w in wf.worksheets}):
                hits['3 비어 있는 셀을 읽는다'].append(
                    '%-34s %-20.20s -> %s   (%s)' % (key, labels.get(key, ''), tgt, f[:40]))
                break

    # ---------------------------------------------------- 4. selected vs limit
    for ws in wf.worksheets:
        if ws.title in BULK:
            continue
        for row in ws.iter_rows(min_col=6, max_col=6):
            for c in row:
                if not isinstance(c.value, (int, float)):
                    continue
                d = wv[ws.title]['D%d' % c.row].value
                lab = ws['C%d' % c.row].value or ''
                if not isinstance(d, (int, float)) or d == 0:
                    continue
                low = str(lab).lower()
                if 'not a minimum' in low or 'ceil' in low:
                    continue          # explicitly an upper bound, see C.22
                if 'zcd.l' in low:
                    continue          # [95] is an equality, not a bound
                lo = 'min' in low
                hi = 'max' in low
                if hi and c.value > d * 1.001:
                    hits['4 선정값이 자기 상한을 넘는다'].append(
                        '%s!F%d  %-24.24s  선정 %g > 상한 %g' % (ws.title, c.row, lab, c.value, d))
                elif lo and c.value < d * 0.999:
                    hits['4 선정값이 자기 하한 아래'].append(
                        '%s!F%d  %-24.24s  선정 %g < 하한 %g' % (ws.title, c.row, lab, c.value, d))

    # ------------------------------------------------------ 5. buried constants
    NUM = re.compile(r'(?<![A-Z0-9_.$])(\d+\.?\d*)(?![0-9.]*[)]?\s*[%A-Z])')
    known = {}
    for key, v in values.items():
        if isinstance(v, (int, float)) and key not in formulas and abs(v) > 1:
            known.setdefault(round(float(v), 6), key)
    for key, f in formulas.items():
        if key.split('!')[0] in BULK:
            continue
        for m in NUM.finditer(f):
            try:
                x = round(float(m.group(1)), 6)
            except ValueError:
                continue
            if x in known and x > 10 and known[x].split('!')[0] not in BULK:
                hits['5 셀로도 존재하는 값을 수식에 박아 넣음'].append(
                    '%-30s %s  <- %g 는 %s 에도 있다' % (key, f[:48], x, known[x]))
                break

    # ------------------------------------------- 7. a selection that goes nowhere
    # The workbook's idiom is IF(F{r}, F{r}, D{r}): a selected part overrides the
    # calculated one. A consumer that reads D{r} directly ignores the selection,
    # and that stays invisible for as long as the two happen to agree.
    SELS = {}
    for ws in wf.worksheets:
        if ws.title in BULK:
            continue
        for row in ws.iter_rows(min_col=6, max_col=6):
            for c in row:
                fv = wv[ws.title].cell(c.row, 6).value
                dv = wv[ws.title].cell(c.row, 4).value
                if isinstance(fv, (int, float)) and isinstance(dv, (int, float)):
                    SELS[(ws.title, c.row)] = (fv, dv, ws.cell(c.row, 3).value)
    GUARD = r"(?:IF|ISNUMBER)\(\s*(?:'[^']+'!)?\$?F\$?%d(?![0-9])"
    SAME = r'(?<![A-Z$!])\$?D\$?%d(?![0-9])'
    OTHER = r"'%s'!\$?D\$?%d(?![0-9])"
    seen = set()
    for key, f in formulas.items():
        sheet, ref = key.split('!')
        for (sh, r), (fv, dv, nm) in SELS.items():
            pat = SAME % r if sh == sheet else OTHER % (re.escape(sh), r)
            if not re.search(pat, f):
                continue
            if re.search(GUARD % r, f):
                continue
            if ref[0] == 'D' and ref[1:].isdigit() and abs(int(ref[1:]) - r) <= 1:
                continue                  # the limit's own ratio cell next to it
            k2 = (key, sh, r)
            if k2 in seen:
                continue
            seen.add(k2)
            gap = abs(fv - dv) / max(abs(dv), 1e-30) * 100
            hits['7 선정값을 무시하고 계산값을 읽는다'].append(
                '%-28s -> %s!D%-4d %-16.16s 선정 %-9g 계산 %-11.6g%s'
                % (key, sh, r, str(nm), fv, dv,
                   '  ** %.1f %% 차이' % gap if gap > 0.3 else ''))

    # ----------------------------------------------------- 6. label vs formula
    for key, f in formulas.items():
        lab = str(labels.get(key, '')).lower()
        if key.split('!')[0] in BULK or not lab:
            continue
        if re.search(r'\bmax\b|\.max|max\.', lab) and 'MIN(' in f and 'MAX(' not in f:
            hits['6 라벨은 max 인데 수식은 MIN'].append(
                '%-30s %-24.24s %s' % (key, lab, f[:50]))
        if re.search(r'\bmin\b|\.min|min\.', lab) and 'MAX(' in f and 'MIN(' not in f:
            hits['6 라벨은 min 인데 수식은 MAX'].append(
                '%-30s %-24.24s %s' % (key, lab, f[:50]))

    LOW = {'5 셀로도 존재하는 값을 수식에 박아 넣음'}
    total = 0
    for cat in sorted(hits):
        if cat in LOW and not verbose:
            print('%s : %d건  (-v 로 보기)' % (cat, len(hits[cat])))
            continue
        print('■ %s — %d건' % (cat, len(hits[cat])))
        for x in hits[cat][:40]:
            print('   ' + x)
        if len(hits[cat]) > 40:
            print('   ... 그 외 %d건' % (len(hits[cat]) - 40))
        print()
        total += len(hits[cat])
    print('설계 시트에서 살펴볼 곳 %d건 (bulk 시트 %s 제외)' % (total, sorted(BULK)))


if __name__ == '__main__':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    main()
