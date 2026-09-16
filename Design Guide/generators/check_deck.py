# -*- coding: utf-8 -*-
"""Geometry check for the training deck - the deck's equivalent of check_sm.py.

There is no renderer on this machine (PowerPoint is installed but COM
automation fails here, the same way it does for Excel), so a slide cannot be
looked at after editing. This checks the things that go wrong when you place
shapes by arithmetic:

  off-slide     a shape that runs past the edge
  overlap       two shapes on top of each other - the complaint that started
                this work was an annotation printed over a legend
  squashed      a picture whose aspect ratio no longer matches its file, which
                is what silently distorts a chart
  overflow      more text than the box can hold at that font size
  furniture     a slide that lost its title, kicker, logo or footer

    python check_deck.py            every slide
    python check_deck.py 17 26      just these
"""
import io
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8',
                              errors='replace')

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
TRAIN = os.path.join(ROOT, 'Training Material')

EMU = 914400.0
W, H = 10.0, 5.625
# a shape may sit this far outside before it counts as off-slide
EDGE = 0.02
# two shapes may share this much area before it counts as an overlap
OVER = 0.045
# rough Arial metrics: characters per inch of line, and line height per point
CPI = 13.5         # Arial: average glyph is about half the point size
LINE = 1.30         # line height as a multiple of the font size


def box(sh):
    return (sh.left / EMU, sh.top / EMU,
            (sh.left + sh.width) / EMU, (sh.top + sh.height) / EMU)


def area(a, b):
    dx = min(a[2], b[2]) - max(a[0], b[0])
    dy = min(a[3], b[3]) - max(a[1], b[1])
    return dx * dy if dx > 0 and dy > 0 else 0.0


def text_of(sh):
    if not sh.has_text_frame:
        return ''
    return sh.text_frame.text


def est_lines(sh):
    """how many lines this text needs in this box, at its own font size"""
    tf = sh.text_frame
    total = 0.0
    height = 0.0
    w = sh.width / EMU - 0.20
    for p in tf.paragraphs:
        t = p.text
        size = 11.0
        for r in p.runs:
            if r.font.size:
                size = r.font.size.pt
                break
        cpl = max(4.0, CPI * (10.0 / size) * w)
        n = max(1, int(len(t) / cpl) + (1 if len(t) % cpl else 0))
        total += n
        height += n * size * LINE / 72.0
    return total, height


def check_slide(prs, no, verbose=False):
    s = prs.slides[no - 1]
    bad = []
    shapes = list(s.shapes)

    furn = {'title': False, 'kicker': False, 'logo': False, 'foot': False}
    for sh in shapes:
        b = box(sh)
        if abs(b[0] - 3.00) < .06 and abs(b[1] - 0.16) < .06:
            furn['title'] = True
        if abs(b[0] - 0.45) < .06 and abs(b[1] - 0.78) < .06:
            furn['kicker'] = True
        if abs(b[0] - 0.30) < .06 and abs(b[1] - 5.00) < .06:
            furn['logo'] = True
        if abs(b[0] - 6.60) < .06 and abs(b[1] - 5.33) < .06:
            furn['foot'] = True

        if b[0] < -EDGE or b[1] < -EDGE or b[2] > W + EDGE or b[3] > H + EDGE:
            bad.append('off-slide  %-11s x %.2f..%.2f  y %.2f..%.2f  %s'
                       % (str(sh.shape_type).split(' ')[0][:11],
                          b[0], b[2], b[1], b[3], text_of(sh)[:26]))

        if sh.shape_type == 13 and hasattr(sh, 'image'):     # PICTURE
            try:
                iw, ih = sh.image.size
                want = iw / ih
                got = (sh.width / sh.height)
                if want and abs(got - want) / want > 0.04:
                    bad.append('squashed   picture  aspect %.3f, file %.3f  '
                               '(%+.0f %%)' % (got, want,
                                               (got / want - 1) * 100))
            except Exception:
                pass

        furniture_box = any(
            abs(b[0] - fx) < .06 and abs(b[1] - fy) < .06
            for fx, fy in ((3.00, 0.16), (0.45, 0.78), (6.60, 5.33),
                           (1.00, 5.33)))
        if sh.has_text_frame and text_of(sh).strip() and not furniture_box:
            n, need = est_lines(sh)
            have = sh.height / EMU
            if need > have * 1.22:
                bad.append('overflow   "%s"  needs ~%.2f in, box %.2f in'
                           % (text_of(sh).replace('\n', ' ')[:34], need, have))

    # overlaps, ignoring the shell (the kicker bar legitimately spans the page)
    body = [sh for sh in shapes
            if not (abs(box(sh)[1] - 0.78) < .06 or abs(box(sh)[1] - 5.33) < .06
                    or abs(box(sh)[1] - 0.16) < .06
                    or abs(box(sh)[1] - 5.00) < .06)]
    for i in range(len(body)):
        for j in range(i + 1, len(body)):
            a, b = box(body[i]), box(body[j])
            ov = area(a, b)
            if ov <= OVER:
                continue
            # a text box sitting inside a card is intentional
            small = min((a[2] - a[0]) * (a[3] - a[1]),
                        (b[2] - b[0]) * (b[3] - b[1]))
            if ov > 0.92 * small:
                continue
            bad.append('overlap    %.2f in²  "%s"  x  "%s"'
                       % (ov, text_of(body[i])[:20] or '(picture)',
                          text_of(body[j])[:20] or '(picture)'))

    # the kicker is optional; the rest is not
    missing = [k for k, v in furn.items() if not v and k != 'kicker']
    if missing and no not in SECTION:
        bad.append('furniture  missing: %s' % ' '.join(missing))
    return bad


# section dividers legitimately have no kicker or title box
SECTION = {1, 8, 18, 31, 35, 40, 48, 55}


def main():
    path = os.path.join(TRAIN, [f for f in os.listdir(TRAIN)
                                if f.lower().endswith('.pptx')][0])
    prs = Presentation(path)
    want = [int(a) for a in sys.argv[1:]] or range(1, len(prs.slides) + 1)
    total = 0
    for n in want:
        bad = check_slide(prs, n)
        if bad:
            total += len(bad)
            print('slide %d' % n)
            for b in bad:
                print('   ' + b)
    print()
    print('%s   슬라이드 %d장 검사 · 문제 %d건'
          % (os.path.basename(path), len(list(want)), total))
    print('VERDICT ' + ('OK' if total == 0 else 'CHECK'))
    return 1 if total else 0


if __name__ == '__main__':
    raise SystemExit(main())
