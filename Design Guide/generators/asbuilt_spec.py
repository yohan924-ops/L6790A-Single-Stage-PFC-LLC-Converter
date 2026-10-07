# -*- coding: utf-8 -*-
"""As-built spec of the three-core transformer the vendor wound (one page). Pictures
from asbuilt_spec_pics.py; measurement values read from the frozen 3-core spec."""
import math, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XImg
import os
S = os.path.dirname(os.path.abspath(__file__)) + '/'          # pictures from asbuilt_spec_pics.py
os.chdir(S + '../../')
# design requirements read from the frozen spec (not typed)
src = openpyxl.load_workbook('Calculation Excel Sheet/variants/Transformer_Spec_7p5to1.xlsx', data_only=True).worksheets[0]
el = {}; fsw = None
for row in src.iter_rows(values_only=True):
    v = [str(x).strip() for x in row if x is not None]
    if len(v) >= 4 and v[1] in ('INDUCTANCE', 'LEAKAGE INDUCTANCE'): el[v[1]] = v[3]
    if len(v) > 2 and v[1].startswith('Switching frequency'): fsw = v[2]
J = 4.5                                   # A/mm2, the density the design guide uses
a = math.pi / 4 * 0.1 ** 2                # one 0.1 mm strand
cap = {'NP1': 130 * a * J, 'NS': 110 * a * J, 'NAUX': math.pi / 4 * 0.4 ** 2 * J}
area = {'NP1': 130 * a, 'NS': 110 * a, 'NAUX': math.pi / 4 * 0.4 ** 2}

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Transformer Spec'
NAVY = '1F3A5F'; LT = 'EEF3F8'
hdr = Font(bold=True, color='FFFFFF', size=11); h1 = Font(bold=True, size=15, color=NAVY)
bold = Font(bold=True); it = Font(italic=True, color='555555')
fill = PatternFill('solid', fgColor=NAVY); fill2 = PatternFill('solid', fgColor=LT)
thin = Side(style='thin', color='BBBBBB'); box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical='center')
for col, wd in zip('ABCDEF', (5, 26, 16, 10, 30, 28)): ws.column_dimensions[col].width = wd
def put(row, vals, font=None, fl=None):
    for i, v in enumerate(vals):
        c = ws.cell(row=row, column=i + 1, value=v); c.alignment = wrap; c.border = box
        if font: c.font = font
        if fl: c.fill = fl
def section(row, title):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    c = ws.cell(row=row, column=1, value=title); c.font = hdr; c.fill = fill
r = 1
ws.cell(row=r, column=1, value='TRANSFORMER SPECIFICATION — as built').font = h1; r += 1
ws.cell(row=r, column=1, value='L6790A single-stage PF LLC, 25 V / 26.3 A.  3 units per set (primaries in series, secondaries in parallel).  Values per unit.').font = it; r += 2

section(r, '1.  WINDING'); r += 1
put(r, ['No', 'Winding', 'Terminal', 'Turns', 'Wire', 'Current capability (%.1f A/mm²)' % J], bold, fill2); r += 1
put(r, [1, 'NP1  Primary', '1 – 4', '5 T', 'Litz 0.1 mm × 130  (%.2f mm²)' % area['NP1'], '%.1f A rms' % cap['NP1']]); r += 1
put(r, [2, 'NS2  Secondary A', '5 – 9', '2 T', 'Litz 0.1 mm × 110, bifilar  (%.2f mm²)' % area['NS'], '%.1f A rms' % cap['NS']]); r += 1
put(r, [3, 'NS3  Secondary B', '7 – 11', '2 T', 'Litz 0.1 mm × 110, bifilar  (%.2f mm²)' % area['NS'], '%.1f A rms' % cap['NS']]); r += 1
put(r, [4, 'NAUX  Auxiliary', '2 – 3', '1 T', 'ø0.4 mm  (%.3f mm²)' % area['NAUX'], '%.2f A rms' % cap['NAUX']]); r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
ws.cell(row=r, column=1, value='Centre tap on the PCB: the NS2 end and the NS3 start that add up (section 4). Not joined inside the transformer.').font = it; r += 2

section(r, '2.  ELECTRICAL REQUIREMENTS   (LCR 100 kHz / 1 V, series L_s, per unit)'); r += 1
put(r, ['No', 'Item', 'Terminal', '', 'Requirement', 'Condition'], bold, fill2); r += 1
put(r, [1, 'INDUCTANCE', '1 – 4', '', el['INDUCTANCE'], 'all other windings OPEN']); r += 1
put(r, [2, 'LEAKAGE INDUCTANCE', '1 – 4', '', el['LEAKAGE INDUCTANCE'], '5–9 and 7–11 both SHORT, 2–3 OPEN']); r += 1
put(r, ['2a', 'Leakage, one half', '1 – 4', '', 'record', '5–9 only SHORT, then 7–11 only SHORT']); r += 1
put(r, [3, 'D.C OVERLAP', '1 – 4', '', '≥ 90 % of initial inductance', '16 A dc, room temperature, other windings OPEN']); r += 2

section(r, '3.  CORE   (A_e 114 mm², l_e 58.82 mm, V_e 6706 mm³, C1 0.52 mm⁻¹, ≈ 33.8 g/pair, material: TBD)'); r += 1
put(r, ['', 'A', '42.00 ± 0.60', '', 'E', '6.00 ± 0.20'], None, None); r += 1
put(r, ['', 'B', '7.85 ± 0.15', '', 'F', '4.85 ± 0.15']); r += 1
put(r, ['', 'C', '19.00 ± 0.30', '', 'B − F', '3.00 ± 0.15']); r += 1
put(r, ['', 'D', '36.00 +0.5 / −0.3', '', '', '']); r += 2

section(r, '4.  POLARITY  (before fitting)'); r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
c = ws.cell(row=r, column=1, value=('Drive 1–4 with ~5 V, 50–100 kHz sine. 5–9 and 7–11 read 2/5 of it, 2–3 reads 1/5. '
    'Join one end of NS2 to one end of NS3 and find the pair whose remaining ends read double: those two pins are the centre tap. Near zero = wrong pair.'))
c.alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[r].height = 46; r += 2
put(r, ['', 'Switching frequency', fsw, '', 'at full load', '']); r += 1

# pictures to the right
for name, anchor, w in (('bobbin.png', 'H2', 260), ('core.png', 'H20', 330)):
    im = XImg(S + name); sc = w / im.width; im.width = w; im.height = int(im.height * sc); ws.add_image(im, anchor)
ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
out = 'Calculation Excel Sheet/variants/Transformer_Spec_7p5to1_asbuilt.xlsx'
wb.save(out); print('saved')
