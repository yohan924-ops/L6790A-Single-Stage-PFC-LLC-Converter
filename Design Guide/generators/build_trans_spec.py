# -*- coding: utf-8 -*-
"""벤더에 보낼 트랜스포머 요구사양서를 만든다.  변형을 인자로 받는다.

이것은 우리가 보내는 '사양서' 이지 벤더가 채워 오는 '검사성적서' 가 아니다.
그래서 빈 칸을 하나도 두지 않는다 - 벤더는 자기 양식으로 승인원을 낸다.

값은 변형 시트의 세트 값을 유닛으로 환산한 것이다: 1차가 직렬이라
인덕턴스는 /N.x, 2차가 병렬이라 전류는 /N.x, 권선비는 /N.x.  N.x 는
시트에서 읽는다 - 세 설계점은 3 이고 8to1 은 2 다.
L.L2(2차 누설)는 지그 부유 인덕턴스 수준이라 주지 않는다 - 개방·단락 두
인덕턴스와 권선비가 있으면 T-모델이 완전히 결정된다.

권선비는 적어야 한다. 저턴수 2차의 자기 인덕턴스에는 전선 고리 성분이 섞여
권선비를 담지 못하고(기존품 26OP-LM83W: 계산 2.12 uH, 사양 2.80 uH, 벤더가
+-50 % 를 걸어 놨다), 벤더의 표준 측정 항목 어디에도 권선비 정보가 없다.

핀 배정은 승인원 방식이다 - NS2 = 3->5, NS3 = 4->6 (엇갈림).  1턴 동판은
보빈 한쪽으로 들어가 반대쪽으로 나오므로 두 끝이 반대편 핀 뱅크에 떨어진다.
인접핀(3-4, 5-6)으로 적었던 판이 있었는데 전기적으로는 같지만 - 어느 쪽이든
4+5 를 이으면 끝이 3·6 인 같은 센터탭 권선이 된다 - 보빈에 자연스럽지 않다.
센터탭(4+5)은 PCB 에서 만든다.

결선도는 넣지 않는다 - 사용자가 직접 그린 것을 쓴다.
이 파일을 실행하면 사양서가 통째로 덮어써진다. 사용자가 손으로 색·배치를
고쳤을 수 있으므로 **요청받았을 때만 실행할 것.**
"""
import math
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# 시트에서 읽는 값(인덕턴스·전류·A.e·주파수)은 여기 적지 않는다.  손으로 적으면
# 시트를 고칠 때마다 조용히 낡는다 - 6:1 사양서가 실제로 그렇게 f.Min 을 90 kHz
# 로 달고 있었다(참값 110 kHz).  아래에는 시트에서 나오지 않는 것만 둔다.
VARIANTS = {
    '9to1':   dict(label='9 : 1'),
    '7p5to1': dict(label='7.5 : 1'),
    '8to1':   dict(label='8 : 1'),
    '6to1':   dict(label='6 : 1'),
}
if len(sys.argv) < 2 or sys.argv[1] not in VARIANTS:
    raise SystemExit('변형을 인자로 주십시오: %s   '
                     '인자 없이 실행해서 정본 사양서를 덮어쓰는 사고가 두 번 났습니다.'
                     % ' / '.join(VARIANTS))
VAR = sys.argv[1]
V = VARIANTS[VAR]

# ---------------------------------------------------- 시트에서 읽어 온다
import check_trans_spec as CTS

_sm = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   '..', '..', 'Smath', 'variants', 'L6790A_%s.sm' % VAR)
if not os.path.exists(_sm):
    raise SystemExit('변형 시트가 없다.  먼저 만들 것:  '
                     'set L6790_VARIANT=%s && python build_l6790_smath.py' % VAR)
_d = CTS.derive(CTS.sheet(_sm))

V.update(
    Np=_d['Np'], Ns=_d['Ns'], Nx=_d['Nx'], Naux=_d['Naux'],
    Lopen=_d['Lopen'], Lshort=_d['Lshort'],
    Ae=int(round(_d['Ae'])),
    Ipri='%.1f A   /   %.1f A' % (_d['Ipri_rms'], _d['Ipri_pk']),
    Isec='%.1f A   /   %.1f A' % (_d['Isec_rms'], _d['Isec_pk']),
    # 개방 시험은 2차 기자력 상쇄가 없으므로 같은 전류로도 자속이 훨씬 크다.
    # 동작 자속을 개방 시험에서 재현하는 전류에 여유를 곱한 값이 시험 전류다.
    Isat=int(math.ceil(_d['Isat_spec'])),
    fsw='%d  to  %d kHz' % (round(_d['fmin']), round(_d['fmax'])),
)

# 공차는 개방·단락 모두 +-10 % 로 통일한다.  버틸 수 있는 개방 인덕턴스 저하는
# 1 - 1/k.floor^2 이고 변형마다 다르지만(9:1 14.9 % · 7.5:1 4.1 % · 6:1 58.9 %),
# 비대칭 공차는 양산 공정이 맞추지 못한다 - 2026-09-07 사용자 결정.
# 7.5:1 을 채택한다면 R.T 를 10 kohm 으로 내려 k.floor 를 1.117 로 올려야 한다.
V['tol_open'] = V['tol_short'] = '±10 %'

NAVY, BAND, LINE, MUTED = '1F3864', 'EDF0F7', 'B4B4B4', '5A6472'
thin = Side(style='thin', color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = 'Transformer Spec'
ws.sheet_view.showGridLines = False

for k, v in {'A': 2.2, 'B': 4.5, 'C': 22, 'D': 16, 'E': 24, 'F': 42,
             'G': 2.2}.items():
    ws.column_dimensions[k].width = v

LAST = 'F'


def put(ref, val, bold=False, size=10, color='000000', fill=None,
        align='left', wrap=False, border=False, italic=False):
    c = ws[ref]
    if val is not None:                 # None 이면 셀에 아무것도 쓰지 않는다
        c.value = val                   # (빈 문자열은 숫자 서식 칸에서 0.00 이 된다)
    c.font = Font(name='Calibri', size=size, bold=bold, italic=italic, color=color)
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
    if fill:
        c.fill = PatternFill('solid', fgColor=fill)
    if border:
        c.border = box
    return c


def span(row):
    ws.merge_cells('B%d:%s%d' % (row, LAST, row))


def section(row, text, tail=None):
    span(row)
    put('B%d' % row, text, bold=True, size=11, color='FFFFFF', fill=NAVY)
    for c in 'CDEF':
        ws['%s%d' % (c, row)].fill = PatternFill('solid', fgColor=NAVY)
    ws.row_dimensions[row].height = 19
    if tail:
        span(row + 1)
        put('B%d' % (row + 1), tail, size=9, italic=True, color=MUTED)
        ws.row_dimensions[row + 1].height = 14


def head(row, cells):
    for col, t in cells:
        put('%s%d' % (col, row), t, bold=True, size=9, fill=BAND,
            align='center', border=True)
    ws.row_dimensions[row].height = 20


def line(row, cells, h=18):
    for col, t, al in cells:
        put('%s%d' % (col, row), t, size=10, align=al, border=True,
            wrap=(al == 'left'))
    ws.row_dimensions[row].height = h


def note(row, text):
    span(row)
    put('B%d' % row, text, size=9, italic=True, color=MUTED)
    ws.row_dimensions[row].height = 14


# ------------------------------------------------------------------ 표제
span(2)
put('B2', 'TRANSFORMER  SPECIFICATION', bold=True, size=18, color=NAVY)
ws.row_dimensions[2].height = 26
span(3)
put('B3', 'L6790A single-stage PF LLC converter      25 V / 26.3 A output'
          '      —      turns ratio %s' % V['label'], size=11, color=MUTED)
ws.row_dimensions[3].height = 17
for c in 'BCDEF':
    ws['%s4' % c].border = Border(bottom=Side(style='medium', color=NAVY))
ws.row_dimensions[4].height = 4

# ------------------------------------------------------------- 1. 구성
section(6, '1.    CONFIGURATION')
COUNT = {2: 'Two', 3: 'Three', 4: 'Four'}.get(V['Nx'], str(V['Nx']))
for i, (t, red) in enumerate([
        ('%s identical transformers per set — primaries in series, '
         'secondaries in parallel.' % COUNT, False),
        ('EVERY VALUE BELOW IS PER TRANSFORMER.', True)]):
    row = 7 + i
    span(row)
    put('B%d' % row, '•   ' + t, size=10, bold=red,
        color=('C00000' if red else '000000'))
    ws.row_dimensions[row].height = 15

# ---------------------------------------------------------- 2. 권선
section(10, '2.    WINDING')
head(11, [('B', 'No'), ('C', 'Winding'), ('D', 'Terminal'), ('E', 'Turns'),
          ('F', 'Winding current           rms   /   peak')])
WIND = [
    ('1', 'NP1     Primary', '1 – 2', '%d Ts' % V['Np'], V['Ipri']),
    ('2', 'NS2     Secondary A', '3 – 5', '%d T' % V['Ns'], V['Isec']),
    ('3', 'NS3     Secondary B', '4 – 6', '%d T' % V['Ns'], V['Isec']),
    ('4', 'NAUX   Auxiliary (ZCD)', 'a – b', '%d T' % V['Naux'], 'sense only'),
]
for i, (n, des, term, turns, cur) in enumerate(WIND):
    line(12 + i, [('B', n, 'center'), ('C', des, 'left'), ('D', term, 'center'),
                  ('E', turns, 'center'), ('F', cur, 'center')])
    ws['E%d' % (12 + i)].font = Font(name='Calibri', size=10, bold=True)

note(17, 'Centre tap is made on the PCB by joining pins 4 and 5 — do NOT join '
         'them inside the transformer.')

# ------------------------------------------------- 3. 전기 요구사양
section(19, '3.    ELECTRICAL  REQUIREMENTS',
        'Measured with an LCR meter at 100 kHz / 1 V.')
head(21, [('B', 'No'), ('C', 'Item'), ('D', 'Terminal'), ('E', 'Requirement'),
          ('F', 'Condition')])
REQ = [
    ('1', 'INDUCTANCE', '1 – 2', '%.2f µH    %s' % (V['Lopen'], V['tol_open']),
     'All other windings OPEN.   L.mag + L.leak, not L.mag alone.', 18),
    ('2', 'LEAKAGE INDUCTANCE', '1 – 2', '%.2f µH    %s' % (V['Lshort'], V['tol_short']),
     'SECONDARY ALL SHORT (NS2 + NS3).   NAUX open.   '
     'Resonant inductor — a target, not a maximum.', 22),
    ('3', 'D.C OVERLAP', '1 – 2', '≥ 90 % of initial inductance',
     'Test current %d A, normal temperature' % V['Isat'], 18),
]
for i, (n, item, term, req, cond, h) in enumerate(REQ):
    row = 22 + i
    line(row, [('B', n, 'center'), ('C', item, 'left'), ('D', term, 'center'),
               ('E', req, 'center'), ('F', cond, 'left')], h=h)
    ws['E%d' % row].font = Font(name='Calibri', size=10, bold=True)
    ws['C%d' % row].alignment = Alignment(horizontal='left', vertical='center',
                                          wrap_text=True)

note(25, 'Both measured at 1 – 2 of ONE transformer.   '
         '%s in series give a total ratio of %s.' % (COUNT, V['label']))
note(26, 'Item 2 follows the existing production part 26OP-LM83W clause 4-2, '
         '"SECONDARY ALL SHORT" — same vendor, same centre-tapped construction.   '
         'The auxiliary is NOT shorted.')

# ------------------------------------------------------- 4. 동작 조건
section(28, '4.    OPERATING  CONDITIONS')
COND = [
    ('Switching frequency', V['fsw'], 'at full load'),
    ('Core effective area', 'Ae  ≥  %d mm²' % V['Ae'],
     'holds the peak flux density at or below 0.20 T'),
]
for i, (item, val, tail) in enumerate(COND):
    row = 29 + i
    put('B%d' % row, '•', size=10, align='center', color=MUTED)
    ws.merge_cells('C%d:D%d' % (row, row))
    put('C%d' % row, item, size=10)
    put('E%d' % row, val, size=10, bold=True)
    if tail:
        put('F%d' % row, tail, size=9, italic=True, color=MUTED)
    ws.row_dimensions[row].height = 16

note(32, 'Core, bobbin, wire and winding arrangement are the supplier’s choice, '
         'provided section 3 is met.')

# ------------------------------------------------------------- 인쇄
ws.page_setup.orientation = 'portrait'
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = 'A1:G33'
ws.page_margins.left = ws.page_margins.right = 0.3
ws.page_margins.top = ws.page_margins.bottom = 0.4

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   '..', '..', 'Calculation Excel Sheet',
                   'variants', 'Transformer_Spec_%s.xlsx' % VAR)
os.makedirs(os.path.dirname(OUT), exist_ok=True)
try:
    wb.save(OUT)
except PermissionError:
    raise SystemExit('사양서가 Excel 에서 열려 있어 쓸 수 없다.  닫고 다시 실행할 것:  '
                     + os.path.normpath(OUT))
print('생성:', os.path.normpath(OUT))
